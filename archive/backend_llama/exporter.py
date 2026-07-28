from docx import Document
import os

SECTIONS = ['BACKGROUND', 'THEMES', 'INTERVENTION', 'RESULTS', 'LEARNING OUTCOMES']

def export(company_name: str, case_text: str, citation_style: str = 'APA') -> str:
    """Build IFQM-structured Word doc and return file path."""
    doc = Document()
    doc.add_heading(f'{company_name}: Case Study', level=1)

    for section in SECTIONS:
        doc.add_heading(section, level=2)
        start = case_text.find(section + ':')
        if start != -1:
            start += len(section) + 1
            ends = [case_text.find(s + ':', start) for s in SECTIONS if s != section]
            end = min([e for e in ends if e > 0], default=len(case_text))
            doc.add_paragraph(case_text[start:end].strip())

    doc.add_heading('Exhibits', level=2)
    doc.add_paragraph('[Exhibits populated from source documents]')

    doc.add_heading('References', level=2)
    doc.add_paragraph(f'[Citations formatted in {citation_style} style]')

    os.makedirs('../outputs', exist_ok=True)
    path = f'../outputs/{company_name}_case_study.docx'
    doc.save(path)
    return path
